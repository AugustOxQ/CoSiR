#!/usr/bin/env python3
"""Build a weekly slide deck (.pptx) from its slide markdown. Reusable template builder.

Usage:  python build_weekly_slides.py <slides.md> [--out deck.pptx] [--footer "text"]
  --out     default: docs/reports/pptx/<md stem>.pptx under the repo that contains the markdown
            (the nearest parent folder holding docs/reports), else next to the markdown
  --footer  default: the title slide's "# " line
Needs python-pptx and Pillow. The design (16:9, teal accents, Calibri, striped tables, footer, speaker notes,
monospace code box) is the one the user approved on 2026-10-07 for the CoSiR weekly deck.

The markdown format is strict and simple, so edits to it rebuild without touching this script:
slides are separated by a line "---"; the first slide starts with "# " (title slide), every other
slide with "## "; a slide holds at most one image "![alt](path)", tables as "|" rows, bullets as
"- " lines, and a final "Notes:" line that goes to the speaker notes. "**x**" marks bold text.
Image paths are relative to the markdown file. A fenced ```text block is drawn as a monospace code box.
"""
import math
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SRC = None      # set from the command line in main()
OUT = None
FOOTER = ""

W_IN, H_IN = 13.333, 7.5
LEFT, RIGHT = 0.5, 12.833
TOP, BOTTOM = 1.25, 7.0
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x57, 0x60, 0x6A)
TEAL = RGBColor(0x1B, 0x8A, 0x7A)
HEADER_FILL = RGBColor(0xE3, 0xF6, 0xF3)
ROW_FILL = RGBColor(0xFF, 0xFF, 0xFF)
ALT_FILL = RGBColor(0xF7, 0xF8, 0xFA)
FONT = "Calibri"
MONO = "DejaVu Sans Mono"
CODE_FILL = RGBColor(0xEE, 0xEF, 0xF1)
WARNINGS = []


# ---------------------------------------------------------------- parsing
def parse(md: str):
    slides = []
    for block in re.split(r"\n---\n", md.strip()):
        slide = {"title": "", "level": 2, "lines": [], "image": None, "table": [], "bullets": [],
                 "notes": "", "code": []}
        rows = block.strip().splitlines()
        i = 0
        while i < len(rows):
            line = rows[i].rstrip()
            if line.startswith("```"):
                i += 1
                while i < len(rows) and not rows[i].startswith("```"):
                    slide["code"].append(rows[i].rstrip())
                    i += 1
                i += 1
                continue
            if line.startswith("# ") and not slide["title"]:
                slide["title"], slide["level"] = line[2:].strip(), 1
            elif line.startswith("## ") and not slide["title"]:
                slide["title"] = line[3:].strip()
            elif line.startswith("!["):
                m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
                slide["image"] = (SRC.parent / m.group(2)).resolve()
            elif line.startswith("|"):
                slide["table"].append(line)
            elif line.startswith("- "):
                slide["bullets"].append(line[2:].strip())
            elif line.startswith("Notes:"):
                slide["notes"] = " ".join([line[6:].strip()] + [r.strip() for r in rows[i + 1:]]).strip()
                break
            elif line.strip():
                slide["lines"].append(line.strip())
            i += 1
        slides.append(slide)
    return slides


def parse_table(lines):
    cells = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines]
    align = []
    body = []
    for row in cells:
        if all(re.fullmatch(r":?-{3,}:?", c) for c in row):
            align = ["right" if c.endswith(":") and not c.startswith(":") else "left" for c in row]
        else:
            body.append(row)
    return body, align or ["left"] * len(body[0])


def clean(text):
    return text.replace("`", "")


def add_runs(paragraph, text, size, color=INK, bold=False):
    """Add text to a paragraph, honouring **bold** spans."""
    parts = re.split(r"(\*\*[^*]+\*\*)", clean(text))
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        is_bold = part.startswith("**") and part.endswith("**")
        run.text = part[2:-2] if is_bold else part
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold or is_bold
        run.font.color.rgb = color


# ---------------------------------------------------------------- text sizing
def text_height(texts, width_in, size, spacing=1.18, gap_pt=6):
    """Rough height (inches) that a list of paragraphs needs at a given font size."""
    chars_per_line = max(10, width_in * 72 / (size * 0.50))
    lines = sum(max(1, math.ceil(len(clean(t).replace("**", "")) / chars_per_line)) for t in texts)
    return lines * size * spacing / 72 + len(texts) * gap_pt / 72


def fit_size(texts, width_in, height_in, sizes=(22, 20, 18, 17, 16, 15, 14, 13, 12, 11)):
    for s in sizes:
        if text_height(texts, width_in, s) <= height_in:
            return s
    return sizes[-1]


# ---------------------------------------------------------------- shapes
def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    return tf


def add_bullets(slide, bullets, x, y, w, h, size=None, name=""):
    size = size or fit_size(bullets, w - 0.3, h)
    need = text_height(bullets, w - 0.3, size)
    if need > h + 0.05:
        WARNINGS.append(f"{name}: bullets may overflow ({need:.2f} in > {h:.2f} in at {size} pt)")
    tf = textbox(slide, x, y, w, h)
    for k, text in enumerate(bullets):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        p.line_spacing = 1.08
        bullet = p.add_run()
        bullet.text = "•  "
        bullet.font.name, bullet.font.size = FONT, Pt(size)
        bullet.font.color.rgb = TEAL
        bullet.font.bold = True
        add_runs(p, text, size)
    return need


def add_code(slide, lines, x, y, w, h, size=10.5, name=""):
    """Monospace code box with light grey fill."""
    cpl = (w - 0.3) * 72 / (size * 0.60)
    n = sum(max(1, math.ceil(len(l) / cpl)) for l in lines)
    need = n * size * 1.2 / 72 + 0.2
    if need > h + 0.05:
        WARNINGS.append(f"{name}: code block may overflow ({need:.2f} in > {h:.2f} in at {size} pt)")
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(min(max(need + 0.4, 0.5), h)))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_FILL
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    for k, text in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.name, run.font.size = MONO, Pt(size)
        run.font.color.rgb = INK
    return need


def add_table(slide, lines, x, y, w, max_h, name=""):
    body, align = parse_table(lines)
    n_rows, n_cols = len(body), len(body[0])
    size = 16 if n_rows <= 5 else 15 if n_rows <= 7 else 14
    lens = [max(len(clean(r[c]).replace("**", "")) for r in body) for c in range(n_cols)]
    weights = [max(8, min(l, 60)) for l in lens]
    widths = [w * wt / sum(weights) for wt in weights]
    # row heights from wrapped line counts
    row_h = []
    for r in body:
        lines_needed = max(max(1, math.ceil(len(clean(c).replace("**", "")) /
                                          max(6, widths[j] * 72 / (size * 0.52))))
                           for j, c in enumerate(r))
        row_h.append(0.12 + lines_needed * size * 1.2 / 72)
    total = sum(row_h)
    if total > max_h:
        scale = max_h / total
        if scale < 0.85:
            WARNINGS.append(f"{name}: table may be tight ({total:.2f} in > {max_h:.2f} in)")
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w),
                                   Inches(min(total, max_h)))
    table = shape.table
    for j, cw in enumerate(widths):
        table.columns[j].width = Emu(int(Inches(cw)))
    for i, r in enumerate(body):
        table.rows[i].height = Emu(int(Inches(row_h[i])))
        for j, text in enumerate(r):
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEADER_FILL if i == 0 else (ALT_FILL if i % 2 == 0 else ROW_FILL)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (align[j] == "right" and i > 0) else PP_ALIGN.LEFT
            add_runs(p, text, size, bold=(i == 0))
    return min(total, max_h)


def add_image(slide, path, x, y, max_w, max_h, name=""):
    if not path.exists():
        WARNINGS.append(f"{name}: missing image {path}")
        return 0, 0
    with Image.open(path) as im:
        aspect = im.width / im.height
    w = max_w
    h = w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    slide.shapes.add_picture(str(path), Inches(x + (max_w - w) / 2), Inches(y), Inches(w), Inches(h))
    return w, h


def image_aspect(path):
    with Image.open(path) as im:
        return im.width / im.height


def chrome(slide, title, number):
    tf = textbox(slide, LEFT, 0.32, RIGHT - LEFT, 0.7, anchor=MSO_ANCHOR.MIDDLE)
    add_runs(tf.paragraphs[0], title, 26, bold=True)
    bar = slide.shapes.add_shape(1, Inches(LEFT), Inches(1.06), Inches(1.4), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    ft = textbox(slide, LEFT, 7.08, 8, 0.3)
    add_runs(ft.paragraphs[0], FOOTER, 10, color=MUTED)
    nt = textbox(slide, RIGHT - 1.0, 7.08, 1.0, 0.3)
    nt.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_runs(nt.paragraphs[0], str(number), 10, color=MUTED)


# ---------------------------------------------------------------- layouts
def title_slide(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, 0, Inches(2.35), Inches(W_IN), Inches(0.06))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    tf = textbox(slide, 0.9, 2.6, 11.5, 1.2, anchor=MSO_ANCHOR.MIDDLE)
    add_runs(tf.paragraphs[0], s["title"], 38, bold=True)
    sub = textbox(slide, 0.9, 3.9, 11.5, 1.8)
    for k, line in enumerate(s["lines"]):
        p = sub.paragraphs[0] if k == 0 else sub.add_paragraph()
        p.space_after = Pt(8)
        add_runs(p, line, 20 if k == 0 else 16, color=INK if k == 0 else MUTED)
    return slide


def content_slide(prs, s, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, s["title"], number)
    name = f"slide {number}"
    avail_h = BOTTOM - TOP
    width = RIGHT - LEFT
    img, table, bullets = s["image"], s["table"], s["bullets"]
    if img and table:
        # figure on the left, table on the right
        _, ih = add_image(slide, img, LEFT, TOP, 7.9, avail_h, name)
        add_table(slide, table, LEFT + 8.1, TOP + 0.2, width - 8.1, avail_h - 0.2, name)
    elif img:
        aspect = image_aspect(img)
        if aspect >= 1.95:
            size = 16
            bh = text_height(bullets, width - 0.3, size) if bullets else 0
            ih = min(width / aspect, avail_h - bh - 0.15)
            _, ih = add_image(slide, img, LEFT, TOP, width, ih, name)
            if bullets:
                add_bullets(slide, bullets, LEFT, TOP + ih + 0.15, width, avail_h - ih - 0.15, size, name)
        else:
            iw = min(8.3, avail_h * aspect)
            iw_used, _ = add_image(slide, img, LEFT, TOP, iw, avail_h, name)
            bx = LEFT + iw_used + 0.35
            add_bullets(slide, bullets, bx, TOP + 0.1, RIGHT - bx, avail_h - 0.1, None, name)
    elif table:
        bh = text_height(bullets, width - 0.3, 18) if bullets else 0
        th = add_table(slide, table, LEFT, TOP + 0.1, width, avail_h - bh - 0.35, name)
        if bullets:
            add_bullets(slide, bullets, LEFT, TOP + 0.1 + th + 0.25, width, avail_h - th - 0.35, 18, name)
    elif s["code"]:
        size = 16
        bh = text_height(bullets, width - 0.3, size)
        add_bullets(slide, bullets, LEFT, TOP + 0.1, width, bh + 0.1, size, name)
        cy = TOP + 0.1 + bh + 0.25
        add_code(slide, s["code"], LEFT, cy, width, BOTTOM + 0.05 - cy, 10.5, name)
    else:
        add_bullets(slide, bullets, LEFT, TOP + 0.2, width, avail_h - 0.2, None, name)
    return slide


def _default_out(src):
    for parent in src.resolve().parents:
        if (parent / "docs" / "reports").is_dir():
            return parent / "docs" / "reports" / "pptx" / f"{src.stem}.pptx"
    return src.with_suffix(".pptx")


def main():
    global SRC, OUT, FOOTER
    import argparse
    ap = argparse.ArgumentParser(description="Build a .pptx deck from a slide markdown file.")
    ap.add_argument("slides_md")
    ap.add_argument("--out")
    ap.add_argument("--footer")
    args = ap.parse_args()
    SRC = Path(args.slides_md).resolve()
    OUT = Path(args.out).resolve() if args.out else _default_out(SRC)
    text = SRC.read_text(encoding="utf-8")
    m = re.search(r"^# (.+)$", text, re.M)
    FOOTER = args.footer if args.footer is not None else (m.group(1).strip() if m else "")
    slides = parse(text)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    for k, s in enumerate(slides, start=1):
        slide = title_slide(prs, s) if s["level"] == 1 else content_slide(prs, s, k)
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT} ({len(slides)} slides)")
    for w in WARNINGS:
        print("WARNING", w)
    return 1 if any("missing image" in w for w in WARNINGS) else 0


if __name__ == "__main__":
    sys.exit(main())
